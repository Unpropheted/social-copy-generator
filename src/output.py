from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt


OUTPUT_DIR = Path("outputs")


def save_copy(copy, image_path):
    image_path = Path(image_path)

    OUTPUT_DIR.mkdir(exist_ok=True)

    output_path = OUTPUT_DIR / f"{image_path.stem}_copy.txt"

    output_path.write_text(
        copy,
        encoding="utf-8"
    )

    return output_path


def save_copies(results):
    OUTPUT_DIR.mkdir(exist_ok=True)

    timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")

    output_path = OUTPUT_DIR / f"copys_{timestamp}.docx"

    document = Document()

    # Estilo general del documento
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(11)

    for site_index, (site_name, posts) in enumerate(
        _group_results_by_site(results)
    ):

        if site_index > 0:
            document.add_page_break()

        # Nombre del sitio
        document.add_heading(site_name, level=1)

        for post_index, (image_paths, copy) in enumerate(
            posts,
            start=1
        ):

            document.add_heading(
                f"Publicación {post_index}",
                level=2
            )

            if len(image_paths) == 1:
                document.add_paragraph(
                    "Tipo: Publicación individual"
                )

                document.add_paragraph(
                    f"Imagen: {image_paths[0].name}"
                )

            else:
                document.add_paragraph(
                    "Tipo: Carrusel"
                )

                document.add_paragraph(
                    "Imágenes:"
                )

                for image_path in image_paths:
                    document.add_paragraph(
                        image_path.name,
                        style="List Bullet"
                    )

            document.add_paragraph(
                "COPY:",
                style=None
            )

            document.add_paragraph(
                copy.strip()
            )

    document.save(output_path)

    return output_path


def _group_results_by_site(results):
    sites = {}

    for site_name, image_paths, copy in results:
        if site_name not in sites:
            sites[site_name] = []

        sites[site_name].append(
            (image_paths, copy)
        )

    return sites.items()